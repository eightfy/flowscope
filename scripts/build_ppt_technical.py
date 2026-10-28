"""Build the 9-slide, 5-minute technical FlowScope report deck."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ppt_common import (
    ACCENT,
    DIM,
    EA_FONT,
    INK,
    LATIN_FONT,
    LINE,
    M,
    MUTED,
    OUT as _OLD_OUT,
    SH,
    SW,
    WHITE,
    add_line,
    add_notes,
    add_paragraphs,
    add_rect,
    add_table,
    add_text,
    header as _old_header,
    load_data,
    make_chart,
)

OUT = Path(__file__).resolve().parent.parent / "FlowScope_5min.pptx"
TOTAL = 9


def header(slide, kicker, title, subtitle=None, num=None):
    add_text(slide, kicker, M, 0.38, 7.5, 0.22, size=10, color=MUTED)
    add_text(slide, title, M, 0.66, 11.6, 0.55, size=27, color=INK, bold=True, line_spacing=1.0)
    if subtitle:
        add_text(slide, subtitle, M, 1.23, 11.4, 0.3, size=10.5, color=MUTED)
    add_line(slide, M, 1.56, SW - M, 1.56, LINE, 0.6)
    if num is not None:
        add_text(slide, f"{num:02d} / {TOTAL:02d}", SW - 1.22, 7.1, 0.65, 0.2, size=10, color=DIM, align=PP_ALIGN.RIGHT)


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    return s


def block(slide, x, y, w, title, bullets, size=10):
    add_line(slide, x, y, x + w, y, LINE, 0.6)
    add_text(slide, title, x, y + 0.12, w, 0.28, size=12.5, color=INK, bold=True)
    add_paragraphs(
        slide,
        [{"runs": [{"text": "·  ", "color": ACCENT, "bold": True}, {"text": t}]} for t in bullets],
        x,
        y + 0.5,
        w,
        1.35,
        size=size,
        color=INK,
        line_spacing=1.32,
        space_after=3,
    )


def build():
    profile, gb, rf = load_data()
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    # 1 cover
    s = new_slide(prs)
    add_text(s, "FLOWSCOPE", 0.87, 0.55, 4.0, 0.25, size=10, color=MUTED)
    add_text(s, "区域流量分析与预测系统", 0.84, 2.05, 11.8, 0.95, size=40, color=INK, bold=True, line_spacing=1.0)
    add_text(s, "技术实现与结果汇报", 0.87, 3.08, 8.0, 0.4, size=15, color=MUTED)
    add_line(s, 0.87, 4.0, 3.62, 4.0, INK, 1.4)
    add_text(s, "项目汇报 · 5 分钟", 0.87, 6.47, 4.0, 0.25, size=10, color=MUTED)
    add_text(s, "2026.10", 10.95, 6.47, 1.55, 0.25, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    add_text(s, "3 个数据集  ·  15 种算法  ·  5 个功能页面", 7.85, 5.92, 4.65, 0.25, size=10, color=DIM, align=PP_ALIGN.RIGHT)
    add_text(s, "01 / 09", 11.95, 7.1, 0.65, 0.2, size=10, color=DIM, align=PP_ALIGN.RIGHT)
    add_notes(s, "开场 15 秒：这是一个区域流量分析与预测的本地 Web 应用。汇报分成三部分：数据与任务、技术实现、结果与限制。")

    # 2 data and task
    s = new_slide(prs)
    header(s, "DATA & TASK", "数据概况与预测任务", "区域级宽表，按时间顺序切分训练、验证与测试", 2)
    data = [
        ["数据集", "区域", "时间", "粒度", "单位"],
        ["米兰", "88", "2013.11 - 2014.01", "10 min", "流量"],
        ["特伦蒂诺", "223", "2013.11 - 2014.01", "10 min", "流量"],
        ["济南", "148", "2024.07 - 2024.08", "15 min", "字节"],
    ]
    add_table(s, data, 0.83, 1.9, [1.25, 0.75, 2.95, 0.95, 1.1], row_h=0.42, font_size=10)
    add_text(s, "宽表：第一列时间，其余列区域；预处理后无缺失。", 0.83, 3.92, 7.0, 0.25, size=10, color=MUTED)
    add_text(s, "8,928 / 8,928 / 2,784 个时间点", 0.83, 4.42, 4.8, 0.3, size=12.5, color=INK, bold=True)
    add_text(s, "20,640 个观测点 · 459 个区域", 0.83, 4.76, 4.8, 0.25, size=10, color=MUTED)

    add_text(s, "TASK", 8.55, 1.86, 2.0, 0.22, size=10, color=MUTED)
    add_text(s, "区域级多步预测", 8.53, 2.16, 4.0, 0.35, size=17, color=INK, bold=True)
    add_text(s, "x[t-L+1 : t]  →  x[t+1 : t+H]", 8.53, 2.72, 4.2, 0.3, size=13, color=ACCENT)
    add_paragraphs(
        s,
        [
            {"runs": [{"text": "L", "bold": True}, {"text": "：输入回看步数"}, {"text": "   H", "bold": True}, {"text": "：预测步数"}]},
            {"runs": [{"text": "N", "bold": True}, {"text": "：训练历史长度，三者均可调"}]},
            {"runs": [{"text": "验证集：训练数据末尾 18%，按时间顺序切分"}]},
            {"runs": [{"text": "基线：季节朴素（season = 一天步数）"}]},
        ],
        8.55,
        3.22,
        4.0,
        1.55,
        size=10,
        color=INK,
        line_spacing=1.4,
        space_after=3,
    )

    add_line(s, 0.83, 5.42, 12.5, 5.42, LINE, 0.6)
    add_text(s, "预处理", 0.83, 5.62, 1.5, 0.22, size=10, color=MUTED)
    add_text(s, "CSV → T×N float32 NPZ；区域与 GeoJSON 按名称求交集；缺失值前后填充。", 0.83, 5.88, 7.2, 0.45, size=10, color=MUTED)
    add_text(s, "区域匹配", 8.65, 5.62, 1.5, 0.22, size=10, color=MUTED)
    add_text(s, "济南 148 / 162 个 GeoJSON 区域有流量；其余区域丢弃。", 8.65, 5.88, 3.8, 0.45, size=10, color=MUTED)
    add_text(s, "数据来源：米兰 / 特伦蒂诺为区域网络流量；济南为 4G 下行字节。", 0.83, 6.62, 8.5, 0.25, size=10, color=DIM)
    add_notes(s, "35 秒：三个数据集都是区域级宽表。米兰 88 个区域、特伦蒂诺 223 个、济南 148 个，时间粒度分别是 10 分钟和 15 分钟。预测任务是给定 L 步历史，输出 H 步未来值，训练、验证、测试严格按时间切分。")

    # 3 architecture & stack
    s = new_slide(prs)
    header(s, "ARCHITECTURE", "系统架构与技术栈", "浏览器端、服务端、计算与数据层", 3)
    layers = [
        ("交互层", "React 19 / Vite 8 / ECharts 6 / 自定义 CSS", "五个页面；地图、折线、热力、散点均使用 ECharts"),
        ("服务层", "FastAPI / Uvicorn / Pydantic / ThreadPoolExecutor(2)", "REST API、静态资源托管、后台训练任务与状态查询"),
        ("计算与数据层", "NumPy / pandas / scikit-learn / statsmodels / SciPy", "NPZ 矩阵、GeoJSON 边界、SQLite 实验档案"),
    ]
    for i, (label, tech, desc) in enumerate(layers):
        y = 1.95 + i * 1.35
        add_line(s, 0.83, y, 8.05, y, LINE, 0.6)
        add_text(s, label, 0.83, y + 0.15, 1.25, 0.25, size=10, color=MUTED)
        add_text(s, tech, 2.05, y + 0.12, 5.8, 0.3, size=12.5, color=INK, bold=True)
        add_text(s, desc, 2.05, y + 0.5, 5.8, 0.35, size=10, color=MUTED)

    add_line(s, 8.55, 1.95, 12.5, 1.95, LINE, 0.6)
    add_text(s, "主要接口", 8.55, 2.08, 2.0, 0.22, size=10, color=MUTED)
    endpoints = [
        "GET  /api/datasets",
        "GET  /api/datasets/{id}/temporal",
        "POST /api/cluster",
        "POST /api/forecast",
        "GET  /api/jobs/{id}",
        "GET  /api/runs",
    ]
    add_paragraphs(s, endpoints, 8.55, 2.4, 3.9, 1.7, size=10, color=INK, line_spacing=1.5, space_after=0)
    add_line(s, 8.55, 4.35, 12.5, 4.35, LINE, 0.6)
    add_text(s, "部署与数据流", 8.55, 4.48, 2.5, 0.22, size=10, color=MUTED)
    add_text(
        s,
        "FastAPI 托管前端 dist；单端口 8000。地图使用本地 GeoJSON，不依赖外部瓦片服务。",
        8.55,
        4.76,
        3.9,
        0.8,
        size=10,
        color=MUTED,
        line_spacing=1.4,
    )
    add_text(s, "GeoJSON + NPZ  →  DataStore  →  API  →  React / ECharts", 0.83, 6.25, 8.4, 0.25, size=10, color=INK)
    add_text(s, "训练任务：POST /api/forecast 创建 job，前端每 650 ms 轮询状态。", 0.83, 6.58, 10.5, 0.25, size=10, color=MUTED)
    add_notes(s, "35 秒：技术上分三层。前端 React 加 ECharts，后端 FastAPI，计算用 NumPy、pandas、scikit-learn 和 statsmodels，数据落在 NPZ、GeoJSON 和 SQLite。服务端同时托管前端构建产物和 REST API，训练任务走后台线程。")

    # 4 backend data and analysis
    s = new_slide(prs)
    header(s, "BACKEND 1/2", "后端实现：数据与分析", "DataStore、时间序列分析、聚类特征与任务存储", 4)
    block(
        s,
        0.83,
        1.9,
        5.65,
        "DataStore",
        [
            "3 个 NPZ，float32 T×N 矩阵；时间索引与区域元数据同时加载",
            "按小时 / 天聚合结果做内存缓存，接口只做切片与重采样",
            "GeoJSON 区域按名称匹配；缺失值前后填充",
        ],
        size=10,
    )
    block(
        s,
        6.7,
        1.9,
        5.8,
        "时间与空间分析",
        [
            "pandas groupby / resample 生成 24h 节律与星期热力",
            "STL（period=7）、ACF（120 阶）、SciPy 周期图",
            "Moran's I：6 近邻权重 + 199 次置换，输出 p 值与 z 分数",
        ],
        size=10,
    )
    block(
        s,
        0.83,
        4.0,
        5.65,
        "功能区聚类",
        [
            "28 维特征：24h 归一化形态、周末比率、流量水平、CV、峰值强度",
            "StandardScaler → KMeans / GaussianMixture / Agglomerative",
            "轮廓系数选择 K；PCA 前两维用于散点展示与解释",
        ],
        size=10,
    )
    block(
        s,
        6.7,
        4.0,
        5.8,
        "任务与存储",
        [
            "ThreadPoolExecutor(2)，任务状态包含进度、阶段、日志、学习曲线",
            "取消事件在训练循环中检查；失败与异常写回任务状态",
            "SQLite runs.db 保存配置、指标、预测序列与训练曲线",
        ],
        size=10,
    )
    add_text(
        s,
        "接口：GET /api/datasets · POST /api/cluster · POST /api/forecast · GET /api/jobs/{id} · GET /api/runs",
        0.83,
        6.72,
        11.5,
        0.25,
        size=10,
        color=MUTED,
    )
    add_notes(s, "45 秒：后端第一块是数据与分析。DataStore 把 NPZ 读成 T×N 矩阵并缓存聚合结果；分析层用 pandas 做重采样，用 statsmodels 做 STL 和 ACF，用 SciPy 做周期图，空间自相关用 6 近邻加 199 次置换。聚类特征 28 维，标准化后再用 KMeans、GMM 或层次聚类。任务和实验记录放在 SQLite。")

    # 5 forecast engine
    s = new_slide(prs)
    header(s, "BACKEND 2/2", "预测引擎与算法", "递归多步预测、15 种算法、训练任务状态", 5)
    add_text(s, "WINDOW CONSTRUCTION", 0.83, 1.80, 3.2, 0.22, size=10, color=MUTED)
    add_line(s, 0.9, 2.95, 10.4, 2.95, LINE, 0.8)
    add_rect(s, 0.9, 2.89, 7.7, 0.12, fill=LINE, line=None)
    add_rect(s, 7.2, 2.89, 1.4, 0.12, fill=INK, line=None)
    add_rect(s, 8.6, 2.89, 1.6, 0.12, fill=ACCENT, line=None)
    add_line(s, 8.6, 2.76, 8.6, 3.18, INK, 0.8)
    add_text(s, "训练历史 2,016 步", 0.9, 2.42, 3.2, 0.25, size=10, color=INK, bold=True)
    add_text(s, "输入 144 步", 7.2, 3.18, 2.0, 0.25, size=10, color=INK, bold=True)
    add_text(s, "输出 24 步", 8.6, 2.42, 2.0, 0.25, size=10, color=ACCENT, bold=True)
    add_text(s, "监督样本 = 2,016 - 144 = 1,872；验证集取末尾 18%。", 0.9, 3.48, 7.2, 0.25, size=10, color=MUTED)
    add_text(s, "递归预测：每次输出回填输入窗口，直到 H 步；预测值裁剪到训练区间，避免反馈发散。", 0.9, 3.80, 9.5, 0.3, size=10, color=MUTED)

    add_line(s, 0.83, 4.28, 12.5, 4.28, LINE, 0.6)
    add_text(s, "ALGORITHMS  /  15", 0.83, 4.42, 4.0, 0.22, size=10, color=MUTED)
    groups = [
        ("基线", ["Naive", "Seasonal Naive", "Moving Average"], 0.83, 2.0),
        ("统计", ["Exp. Smoothing", "Holt-Winters", "SARIMA"], 3.1, 2.35),
        ("机器学习", ["Ridge / Lasso / KNN", "RF / ExtraTrees", "GBDT / HistGB / SVR"], 5.75, 4.4),
        ("深度学习", ["MLP"], 10.65, 1.6),
    ]
    for tag, names, x, w in groups:
        add_text(s, tag, x, 4.72, w, 0.2, size=10, color=DIM)
        add_paragraphs(s, names, x, 4.96, w, 0.9, size=10, color=INK, line_spacing=1.35, space_after=0)
    add_text(s, "注册表统一描述参数类型、范围与默认值，前端据此生成表单。", 0.83, 5.78, 8.0, 0.25, size=10, color=MUTED)

    add_line(s, 0.83, 6.02, 12.5, 6.02, LINE, 0.6)
    add_text(
        s,
        "训练：log1p 可选；StandardScaler；时间顺序验证；MLP 早停；预测值裁剪到训练区间。",
        0.83,
        6.16,
        11.4,
        0.22,
        size=10,
        color=MUTED,
    )
    add_text(
        s,
        "监控：ThreadPoolExecutor(2)；任务状态包含进度、阶段、日志、损失曲线；完成后写 SQLite。",
        0.83,
        6.44,
        11.4,
        0.22,
        size=10,
        color=MUTED,
    )
    add_text(
        s,
        "指标：MAE / RMSE / MAPE / sMAPE / R² / WAPE / MASE / Skill；90% 残差区间。",
        0.83,
        6.72,
        11.4,
        0.22,
        size=10,
        color=MUTED,
    )
    add_notes(s, "45 秒：预测引擎先构造窗口，训练历史 2016 步，输入 144 步，输出 24 步，监督样本 1872 个。算法共 15 种，分基线、统计、机器学习和深度学习。训练时按时间切分验证集，支持 log1p、标准化、早停和预测裁剪。任务状态、损失曲线和日志通过后台线程返回，完成后写 SQLite。")

    # 6 frontend
    s = new_slide(prs)
    header(s, "FRONTEND", "前端实现：页面与交互", "React 单页应用、ECharts 图表、轮询任务状态", 6)
    block(
        s,
        0.83,
        1.9,
        3.65,
        "页面结构",
        [
            "总览：KPI、趋势、地图、热力",
            "时空探索：时间切片、播放、ACF、STL",
            "功能聚类：特征开关、K 选择、PCA",
            "预测实验室：窗口/参数表单、训练监控",
            "评估对比：实验列表、指标表、曲线叠加",
        ],
        size=10,
    )
    block(
        s,
        4.83,
        1.9,
        3.65,
        "状态与请求",
        [
            "dataset context 统一当前数据源",
            "useEffect 拉取分析、聚类与实验数据",
            "训练中每 650 ms 轮询 /api/jobs/{id}",
            "loading / error / empty 状态统一处理",
        ],
        size=10,
    )
    block(
        s,
        8.83,
        1.9,
        3.65,
        "可视化与表单",
        [
            "ECharts：折线、热力、散点、地图",
            "GeoJSON 通过 registerMap 本地注册",
            "算法参数由 registry 元数据生成",
            "CSS Grid + 响应式布局，无组件库",
        ],
        size=10,
    )
    add_line(s, 0.83, 5.55, 12.5, 5.55, LINE, 0.6)
    add_text(
        s,
        "构建：Vite 输出 dist，FastAPI 静态托管；地图不依赖外部瓦片服务。",
        0.83,
        5.75,
        11.4,
        0.25,
        size=10,
        color=MUTED,
    )
    add_text(
        s,
        "交互：区域搜索、时间滑杆与播放、聚类 K 值、算法参数、实验多选对比。",
        0.83,
        6.1,
        11.4,
        0.25,
        size=10,
        color=MUTED,
    )
    add_text(
        s,
        "结果：每个预测实验自动保存，评估页可回读完整预测序列与指标。",
        0.83,
        6.45,
        11.4,
        0.25,
        size=10,
        color=MUTED,
    )
    add_notes(s, "35 秒：前端是 React 单页应用，五个页面共享 dataset context。图表全部用 ECharts，地图用本地 GeoJSON 注册。预测实验室的参数表单由后端算法注册表生成，训练时每 650 毫秒轮询任务状态。构建产物由 FastAPI 托管。")

    # 7 analysis & clustering result
    s = new_slide(prs)
    header(s, "ANALYSIS", "探索性分析与功能区聚类结果", "米兰，88 个区域，62 天；K = 4", 7)
    add_text(s, "24 小时节律（全域均值 = 100）", 0.83, 1.7, 5.5, 0.22, size=10, color=MUTED)
    cats = [str(h) if h % 3 == 0 else "" for h in range(24)]
    make_chart(
        s,
        0.83,
        1.95,
        6.85,
        3.45,
        cats,
        [
            ("均值", [100] * 24, LINE, True),
            ("相对强度", profile, ACCENT, False),
        ],
        y_min=40,
        y_max=140,
        y_unit=20,
    )
    add_text(s, "13:00 · 133", 4.72, 2.02, 1.25, 0.2, size=10, color=ACCENT, bold=True)

    add_line(s, 0.83, 5.55, 7.65, 5.55, LINE, 0.6)
    add_text(s, "13:00 达到平均峰值", 0.83, 5.72, 3.2, 0.25, size=10.5, color=INK, bold=True)
    add_text(s, "夜间仍保持约 60% 的相对强度", 0.83, 5.99, 3.2, 0.22, size=10, color=MUTED)
    add_text(s, "Moran's I = 0.42", 4.2, 5.72, 3.2, 0.25, size=10.5, color=INK, bold=True)
    add_text(s, "6 近邻权重，199 次置换，p = 0.005", 4.2, 5.99, 3.4, 0.22, size=10, color=MUTED)

    add_text(s, "功能区聚类 / K = 4", 7.95, 1.7, 4.5, 0.22, size=10, color=MUTED)
    clusters = [
        ("中负载 · 日间商务型", 9, "10.2%"),
        ("晚间生活型", 50, "56.8%"),
        ("低负载 · 日间商务型", 7, "8.0%"),
        ("中负载 · 日间商务型 B", 22, "25.0%"),
    ]
    yy = 2.05
    for i, (name, size, share) in enumerate(clusters):
        add_text(s, name, 7.95, yy, 3.55, 0.26, size=10.2, color=INK, bold=(i == 1))
        add_text(s, f"{size} 个区域", 11.25, yy, 0.85, 0.22, size=10, color=MUTED)
        add_text(s, share, 12.12, yy, 0.46, 0.22, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
        add_line(s, 7.95, yy + 0.36, 12.58, yy + 0.36, LINE, 0.5)
        yy += 0.62
    add_text(s, "晚间生活型占 56.8%，是米兰最主要的区域形态。", 7.95, 4.72, 4.6, 0.5, size=10, color=MUTED, line_spacing=1.3)

    add_text(
        s,
        "方法：小时聚合；STL 周期 7 天；ACF 120 阶；聚类 28 维特征；轮廓系数 0.319；PCA 前两维解释 83.1%。",
        0.83,
        6.55,
        11.7,
        0.25,
        size=10,
        color=DIM,
    )
    add_notes(s, "40 秒：结果部分。时间上 13 点最高，夜间仍有约 60% 的相对强度。空间上 Moran's I 0.42，p 值 0.005。聚类用 28 维特征，K=4，轮廓系数 0.319，PCA 前两维解释 83.1%。晚间生活型占 56.8%。")

    # 8 prediction result
    s = new_slide(prs)
    header(s, "RESULT", "预测结果与误差分析", "米兰 DUOMO；输入 144 步，输出 24 步，训练历史 2,016 步", 8)
    add_text(s, "10 分钟粒度，预测 4 小时；递归多步、18% 验证、log1p、标准化、预测裁剪。", 0.83, 1.7, 7.0, 0.22, size=10, color=MUTED)
    horizon = gb["horizon"]
    labels = []
    for i, h in enumerate(horizon):
        labels.append(h["time"][11:13] if i % 6 == 0 else "")
    make_chart(
        s,
        0.83,
        1.95,
        7.1,
        3.4,
        labels,
        [
            ("真实值", [h["actual"] for h in horizon], INK, False),
            ("预测值", [h["predicted"] for h in horizon], ACCENT, True),
        ],
        y_min=9000,
        y_max=13500,
        y_unit=1500,
        legend=True,
    )

    metrics = [
        ("MAE", "449", "平均绝对误差"),
        ("MAPE", "4.49%", "平均绝对百分比误差"),
        ("R²", "0.667", "决定系数"),
        ("Skill", "93.6%", "相对季节朴素基线"),
    ]
    for i, (label, value, note) in enumerate(metrics):
        x = 8.15 + (i % 2) * 2.2
        y = 1.92 + (i // 2) * 1.38
        add_text(s, label, x, y, 1.0, 0.2, size=10, color=MUTED)
        add_text(s, value, x, y + 0.2, 2.0, 0.45, size=25, color=INK if i < 2 else ACCENT, bold=True, line_spacing=1.0)
        add_text(s, note, x, y + 0.76, 2.0, 0.2, size=10, color=DIM)

    table_data = [
        ["模型", "MAE", "RMSE", "MAPE", "R²", "Skill"],
        ["GBDT", "449", "578", "4.49%", "0.667", "93.6%"],
        ["随机森林", "2,342", "2,403", "22.59%", "-4.77", "73.5%"],
    ]
    add_table(s, table_data, 0.83, 5.55, [1.9, 1.6, 1.6, 1.6, 1.5, 1.5], row_h=0.34, font_size=10)
    add_text(s, "示例窗口跨元旦，R² 对窗口敏感；平台保留真实结果，不做修饰。", 0.83, 6.58, 8.3, 0.25, size=10, color=MUTED)
    add_notes(s, "45 秒：DUOMO 的 24 步回测。GBDT MAE 449，MAPE 4.49%，相对季节朴素基线 Skill 93.6%；随机森林 MAE 2342，明显更差。这个窗口跨元旦，R² 受节日影响，平台保留真实结果。")

    # 9 conclusion
    s = new_slide(prs)
    header(s, "CLOSING", "结论、限制与后续工作", "实现情况、已知边界与下一步", 9)
    cols = [
        ("01  实现结果", ["3 个数据集、15 种算法、5 个页面", "FastAPI + 后台任务 + SQLite 实验档案", "时空分析、聚类、预测、评估全流程可运行"]),
        ("02  已知限制", ["SARIMA 季节项训练较慢", "预测区间为残差近似，非校准区间", "单机单用户；跨数据集单位不可直接比较"]),
        ("03  后续工作", ["报告导出与实验分享", "实时流数据接入", "多变量与空间相关模型"]),
    ]
    for i, (tag, items) in enumerate(cols):
        x = 0.83 + i * 3.95
        add_text(s, tag, x, 2.0, 3.4, 0.25, size=10, color=MUTED)
        add_paragraphs(
            s,
            [{"runs": [{"text": "·  ", "color": ACCENT, "bold": True}, {"text": t}]} for t in items],
            x,
            2.45,
            3.45,
            2.5,
            size=11.5,
            color=INK,
            line_spacing=1.5,
            space_after=10,
        )
    add_line(s, 0.83, 5.55, 12.5, 5.55, LINE, 0.6)
    add_text(s, "代码与运行说明：README.md；启动脚本：./run.sh；服务默认 127.0.0.1:8000。", 0.83, 5.8, 10.7, 0.3, size=10.5, color=INK)
    add_text(s, "Q & A", 11.3, 5.85, 1.25, 0.3, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    add_notes(s, "25 秒：最后总结。实现结果是三个数据集、15 种算法、五个页面，前后端和实验档案都打通。限制主要是 SARIMA 季节项慢、区间是近似、单机单用户、跨数据集单位不能直接比。后续做导出、实时流和多变量模型。")

    prs.save(OUT)
    print(f"saved {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
