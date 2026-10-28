"""Build a minimal white-background 7-slide FlowScope report deck."""
from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import numpy as np
import pandas as pd
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "FlowScope_5min.pptx"

# ---------- palette ----------
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x6B, 0x72, 0x80)
DIM = RGBColor(0x9C, 0xA3, 0xAF)
LINE = RGBColor(0xE5, 0xE7, 0xEB)
PALE = RGBColor(0xF6, 0xF7, 0xF9)
ACCENT = RGBColor(0x1D, 0x4E, 0xD8)
ACCENT_PALE = RGBColor(0xEE, 0xF3, 0xFF)

EA_FONT = "PingFang SC"
LATIN_FONT = "Arial"
SW, SH = 13.333, 7.5
M = 0.83


# ---------- text helpers ----------
def set_ea(run, name: str = EA_FONT):
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = OxmlElement("a:ea")
        latin = rpr.find(qn("a:latin"))
        if latin is not None:
            latin.addnext(ea)
        else:
            rpr.append(ea)
    ea.set("typeface", name)


def style_run(run, size=14, color=INK, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = LATIN_FONT
    set_ea(run)


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=14,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing=1.12,
    space_after=0,
    wrap=True,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    style_run(r, size=size, color=color, bold=bold)
    return box


def add_paragraphs(
    slide,
    items,
    x,
    y,
    w,
    h,
    size=12,
    color=INK,
    line_spacing=1.2,
    space_after=6,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        if isinstance(item, str):
            item = {"text": item}
        for seg in item.get("runs", [{"text": item.get("text", ""), "size": item.get("size", size), "bold": item.get("bold", False), "color": item.get("color", color)}]):
            r = p.add_run()
            r.text = seg.get("text", "")
            style_run(
                r,
                size=seg.get("size", item.get("size", size)),
                color=seg.get("color", item.get("color", color)),
                bold=seg.get("bold", item.get("bold", False)),
            )
    return box


def add_rect(slide, x, y, w, h, fill=None, line=None, line_width=0.75):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.shadow.inherit = False
    # Remove the theme style reference; LibreOffice otherwise applies its shadow.
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    tf = shape.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return shape


def add_line(slide, x1, y1, x2, y2, color=LINE, width=0.75):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    return conn


def header(slide, kicker, title, subtitle=None, num=None):
    add_text(slide, kicker, M, 0.38, 7.5, 0.22, size=8.5, color=MUTED)
    add_text(slide, title, M, 0.66, 11.6, 0.55, size=28, color=INK, bold=True, line_spacing=1.0)
    if subtitle:
        add_text(slide, subtitle, M, 1.23, 11.2, 0.3, size=11, color=MUTED)
    add_line(slide, M, 1.56, SW - M, 1.56, LINE, 0.6)
    if num is not None:
        add_text(slide, f"{num:02d} / 07", SW - 1.25, 7.1, 0.7, 0.2, size=8, color=DIM, align=PP_ALIGN.RIGHT)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_table(slide, data, x, y, col_widths, row_h=0.38, font_size=10, header_color=MUTED):
    rows, cols = len(data), len(data[0])
    total_w = sum(col_widths)
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(total_w), Inches(row_h * rows))
    table = shape.table
    table.first_row = False
    table.horz_banding = False
    for ci, cw in enumerate(col_widths):
        table.columns[ci].width = Inches(cw)
    for ri in range(rows):
        table.rows[ri].height = Inches(row_h)
        for ci in range(cols):
            cell = table.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.01)
            cell.margin_bottom = Inches(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = str(data[ri][ci])
            style_run(
                r,
                size=font_size,
                color=header_color if ri == 0 else INK,
                bold=(ri == 0),
            )
            # bottom rule only
            tc_pr = cell._tc.get_or_add_tcPr()
            for old in tc_pr.findall(qn("a:lnB")):
                tc_pr.remove(old)
            ln = OxmlElement("a:lnB")
            ln.set("w", str(int(0.5 * 12700)))
            fill = OxmlElement("a:solidFill")
            clr = OxmlElement("a:srgbClr")
            clr.set("val", "E5E7EB")
            fill.append(clr)
            ln.append(fill)
            tc_pr.append(ln)
    return table


# ---------- chart ----------
def make_chart(slide, x, y, w, h, categories, series, y_min=None, y_max=None, y_unit=None, legend=False):
    cd = CategoryChartData()
    cd.categories = categories
    for name, values, color, dash in series:
        cd.add_series(name, values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE, Inches(x), Inches(y), Inches(w), Inches(h), cd
    ).chart
    chart.has_title = False
    chart.font.name = LATIN_FONT
    chart.font.size = Pt(10)
    chart.font.color.rgb = MUTED
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(10)
        chart.legend.font.name = EA_FONT
        chart.legend.font.color.rgb = MUTED
    plot = chart.plots[0]
    plot.vary_by_categories = False
    for i, (_, _, color, dash) in enumerate(series):
        s = chart.series[i]
        s.smooth = False
        s.format.line.color.rgb = color
        s.format.line.width = Pt(2.0 if i == 0 else 1.8)
        if dash:
            s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        s.marker.style = XL_MARKER_STYLE.NONE

    va = chart.value_axis
    ca = chart.category_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.color.rgb = LINE
    va.format.line.width = Pt(0.5)
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minor_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = LINE
    ca.format.line.width = Pt(0.5)
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.has_major_gridlines = False
    for ax in (va, ca):
        ax.tick_labels.font.size = Pt(10)
        ax.tick_labels.font.name = LATIN_FONT
        ax.tick_labels.font.color.rgb = MUTED
    if y_min is not None:
        va.minimum_scale = y_min
    if y_max is not None:
        va.maximum_scale = y_max
    if y_unit:
        va.major_unit = y_unit
    va.tick_labels.number_format = "#,##0"
    va.tick_labels.number_format_is_linked = False
    return chart


# ---------- source data ----------
def load_data():
    ds = np.load(ROOT / "prepared" / "milano.npz", allow_pickle=True)
    times = pd.to_datetime(ds["times"])
    total = pd.Series(ds["values"].sum(axis=1), index=times)
    hour = total.groupby(total.index.hour).mean()
    profile = (hour / hour.mean() * 100).round(1).tolist()

    con = sqlite3.connect(ROOT / "backend" / "runs.db")
    rows = con.execute(
        "select algorithm, result from runs where region_id='DUOMO' order by created_at desc"
    ).fetchall()
    results = {}
    for algo, raw in rows:
        if algo not in results:
            results[algo] = json.loads(raw)
    gb = results.get("gradient_boosting")
    rf = results.get("random_forest")
    if not gb or not rf:
        raise RuntimeError("缺少 DUOMO 的 gradient_boosting / random_forest 实验记录")
    return profile, gb, rf


# ---------- build ----------
def build():
    profile, gb, rf = load_data()
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    def new_slide():
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = WHITE
        return s

    # 1 cover
    s = new_slide()
    add_text(s, "FLOWSCOPE", 0.87, 0.55, 4.0, 0.25, size=9.5, color=MUTED)
    add_text(s, "城市流量智能分析平台", 0.84, 2.05, 11.8, 0.95, size=42, color=INK, bold=True, line_spacing=1.0)
    add_text(s, "从时空探索、功能区识别到多算法预测与评估", 0.87, 3.05, 10.8, 0.45, size=16, color=MUTED)
    add_line(s, 0.87, 4.02, 3.62, 4.02, INK, 1.4)
    add_text(s, "项目汇报 · 5 分钟", 0.87, 6.47, 4.0, 0.25, size=10, color=MUTED)
    add_text(s, "2026.10", 10.95, 6.47, 1.55, 0.25, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    add_text(s, "3 datasets  ·  15 algorithms  ·  5 modules", 7.95, 5.92, 4.55, 0.25, size=8.5, color=DIM, align=PP_ALIGN.RIGHT)
    add_text(s, "01 / 07", 11.95, 7.1, 0.65, 0.2, size=8, color=DIM, align=PP_ALIGN.RIGHT)
    add_notes(s, "开场 15 秒：这是一个把区域流量数据从探索、聚类一直做到多算法预测与评估的本地 Web 应用。今天用 5 分钟讲清三件事：数据是什么、方法怎么搭、结果怎么样。")

    # 2 data
    s = new_slide()
    header(s, "DATA", "三套数据，一个预测问题", "区域级宽表，时间 × 区域", 2)
    data = [
        ["数据集", "区域", "时间跨度", "粒度"],
        ["米兰", "88", "2013.11  -  2014.01", "10 min"],
        ["特伦蒂诺", "223", "2013.11  -  2014.01", "10 min"],
        ["济南", "148", "2024.07  -  2024.08", "15 min"],
    ]
    add_table(s, data, 0.83, 1.9, [1.35, 0.85, 3.55, 1.05], row_h=0.43, font_size=10)
    add_text(s, "宽表：第一列时间，其余列区域；预处理后无缺失。", 0.83, 4.02, 7.1, 0.25, size=9, color=MUTED)
    add_text(s, "8,928  /  8,928  /  2,784 个时间点", 0.83, 4.55, 4.6, 0.25, size=12, color=INK, bold=True)
    add_text(s, "20,640 个观测点 · 459 个区域", 0.83, 4.86, 4.8, 0.25, size=9, color=MUTED)

    add_text(s, "TASK", 8.55, 1.86, 2.0, 0.22, size=8.5, color=MUTED)
    add_text(s, "给定历史，预测未来 H 步", 8.53, 2.2, 4.1, 0.8, size=18, color=INK, bold=True, line_spacing=1.1)
    add_paragraphs(
        s,
        [
            {"runs": [{"text": "输入  ", "color": MUTED}, {"text": "L 步", "bold": True}]},
            {"runs": [{"text": "输出  ", "color": MUTED}, {"text": "H 步", "bold": True}]},
            {"runs": [{"text": "训练历史  ", "color": MUTED}, {"text": "N 步", "bold": True}]},
            {"runs": [{"text": "基线  ", "color": MUTED}, {"text": "季节朴素", "bold": True}]},
        ],
        8.55,
        3.28,
        4.0,
        1.6,
        size=11,
        color=INK,
        line_spacing=1.35,
        space_after=4,
    )
    add_text(s, "窗口长度全部可调", 8.55, 4.9, 4.0, 0.25, size=9, color=ACCENT)

    add_line(s, 0.83, 5.55, 12.5, 5.55, LINE, 0.6)
    add_text(s, "完整时间网格", 0.83, 5.78, 2.8, 0.22, size=8.5, color=MUTED)
    add_text(s, "无缺失；济南 0.18% 零值", 0.83, 6.02, 3.2, 0.3, size=10)
    add_text(s, "区域匹配", 4.45, 5.78, 2.4, 0.22, size=8.5, color=MUTED)
    add_text(s, "GeoJSON 148 / 162 匹配", 4.45, 6.02, 3.3, 0.3, size=10)
    add_text(s, "来源", 8.25, 5.78, 2.4, 0.22, size=8.5, color=MUTED)
    add_text(s, "米兰 / 特伦蒂诺：区域网络流量；济南：4G 下行字节", 8.25, 6.02, 4.2, 0.45, size=8.5, color=MUTED)
    add_notes(s, "30 秒：三套数据都是区域级宽表。米兰和特伦蒂诺是 10 分钟网络流量，济南是 15 分钟 4G 下行字节。任务统一成给定 L 步历史，预测未来 H 步。")

    # 3 architecture
    s = new_slide()
    header(s, "SYSTEM", "五个模块，一条链路", "从数据到评估，单一 Web 入口", 3)
    modules = [
        ("01", "时空探索", "时间节律与空间场"),
        ("02", "功能聚类", "形态特征与分区"),
        ("03", "样本构建", "输入输出窗口"),
        ("04", "多算法预测", "15 种模型与监控"),
        ("05", "评估对比", "指标与多实验"),
    ]
    gap = 0.18
    bw = (SW - 2 * M - 4 * gap) / 5
    for i, (num, name, desc) in enumerate(modules):
        x = M + i * (bw + gap)
        add_rect(s, x, 2.0, bw, 1.55, fill=WHITE, line=LINE)
        add_text(s, num, x + 0.16, 2.16, 0.5, 0.2, size=8.5, color=DIM)
        add_text(s, name, x + 0.16, 2.44, bw - 0.3, 0.3, size=13, color=INK, bold=True)
        add_text(s, desc, x + 0.16, 2.84, bw - 0.3, 0.45, size=9, color=MUTED, line_spacing=1.25)
        if i < 4:
            ax = x + bw + 0.02
            add_text(s, "→", ax, 2.64, 0.14, 0.22, size=11, color=DIM, align=PP_ALIGN.CENTER)

    add_line(s, 0.83, 4.22, 12.5, 4.22, LINE, 0.6)
    layers = [
        ("交互层", "React + ECharts", "五个页面，离线 GeoJSON 地图"),
        ("服务层", "FastAPI + 线程池", "分析、聚类、训练任务与状态"),
        ("数据层", "NPZ + GeoJSON + SQLite", "内存矩阵、区域边界、实验档案"),
    ]
    colw = (SW - 2 * M - 0.6) / 3
    for i, (tag, tech, desc) in enumerate(layers):
        x = M + i * (colw + 0.3)
        add_text(s, tag, x, 4.45, 1.0, 0.22, size=8.5, color=MUTED)
        add_text(s, tech, x, 4.72, colw, 0.32, size=13, color=INK, bold=True)
        add_text(s, desc, x, 5.1, colw, 0.45, size=9, color=MUTED, line_spacing=1.3)
    add_text(s, "本地单页应用，无需外部瓦片服务，打开一个页面即可走完探索、训练与评估。", 0.83, 6.28, 10.8, 0.3, size=10, color=MUTED)
    add_notes(s, "35 秒：系统分五块，前端五个页面，后端是 FastAPI 加线程池任务，数据落在 NPZ、GeoJSON 和 SQLite。地图用本地 GeoJSON，不依赖外部服务。")

    # 4 EDA + clustering
    s = new_slide()
    header(s, "EXPLORE", "看见节律，也看见功能区", "米兰，88 个区域，62 天", 4)
    add_text(s, "24 小时节律（全域均值 = 100）", 0.83, 1.7, 5.5, 0.22, size=8.5, color=MUTED)
    cats = [str(h) if h % 3 == 0 else "" for h in range(24)]
    make_chart(
        s,
        0.83,
        1.95,
        6.85,
        3.45,
        cats,
        [
            ("均值", [100] * 24, LINE, True),
            ("相对强度", profile, ACCENT, False),
        ],
        y_min=40,
        y_max=140,
        y_unit=20,
    )
    add_text(s, "13:00", 5.05, 2.02, 0.45, 0.2, size=8, color=ACCENT, bold=True)
    add_text(s, "峰值 133", 5.05, 2.2, 0.55, 0.2, size=8, color=ACCENT)

    add_line(s, 0.83, 5.55, 7.65, 5.55, LINE, 0.6)
    add_text(s, "13:00 达到平均峰值", 0.83, 5.72, 3.2, 0.25, size=11, color=INK, bold=True)
    add_text(s, "夜间仍保持约 60% 的相对强度", 0.83, 5.99, 3.2, 0.22, size=9, color=MUTED)
    add_text(s, "Moran's I = 0.42", 4.2, 5.72, 3.2, 0.25, size=11, color=INK, bold=True)
    add_text(s, "高负载区域显著集聚，p = 0.005", 4.2, 5.99, 3.3, 0.22, size=9, color=MUTED)

    add_text(s, "功能区聚类 / K = 4", 7.95, 1.7, 4.5, 0.22, size=8.5, color=MUTED)
    clusters = [
        ("中负载 · 日间商务型", 9, "10.2%"),
        ("晚间生活型", 50, "56.8%"),
        ("低负载 · 日间商务型", 7, "8.0%"),
        ("中负载 · 日间商务型 B", 22, "25.0%"),
    ]
    yy = 2.05
    for i, (name, size, share) in enumerate(clusters):
        add_text(s, name, 7.95, yy, 3.55, 0.26, size=10.5, color=INK, bold=(i == 1))
        add_text(s, f"{size} 个区域", 11.55, yy, 1.0, 0.22, size=8.5, color=MUTED)
        add_text(s, share, 12.2, yy, 0.35, 0.22, size=8.5, color=MUTED, align=PP_ALIGN.RIGHT)
        add_line(s, 7.95, yy + 0.36, 12.58, yy + 0.36, LINE, 0.5)
        yy += 0.62
    add_text(s, "晚间生活型占 56.8%，是米兰最主要的区域形态。", 7.95, 4.72, 4.6, 0.5, size=9, color=MUTED, line_spacing=1.3)
    add_text(s, "特征：24h 形态 · 周末比率 · 流量水平 · 波动性", 7.95, 5.45, 4.6, 0.45, size=8.5, color=DIM, line_spacing=1.3)
    add_notes(s, "40 秒：左边是全域 24 小时节律，13 点最高，夜间仍有约 60% 的相对强度。空间上 Moran's I 0.42，高负载区域集聚。右边用 K-Means 把 88 个区域分成 4 类，晚间生活型占 56.8%。")

    # 5 forecast construction + algorithms
    s = new_slide()
    header(s, "FORECAST", "预测：先定窗口，再选算法", "窗口可调，算法可换，训练可看", 5)
    add_text(s, "WINDOWS", 0.83, 1.82, 2.0, 0.22, size=8.5, color=MUTED)
    # timeline
    add_line(s, 0.9, 3.02, 11.15, 3.02, LINE, 0.8)
    add_rect(s, 0.9, 2.96, 6.65, 0.12, fill=LINE, line=None)
    add_rect(s, 7.55, 2.96, 1.8, 0.12, fill=INK, line=None)
    add_rect(s, 9.35, 2.96, 1.8, 0.12, fill=ACCENT, line=None)
    add_line(s, 9.35, 2.82, 9.35, 3.28, INK, 0.8)
    add_text(s, "训练历史 2,016 步", 0.9, 2.48, 3.2, 0.25, size=10, color=INK, bold=True)
    add_text(s, "输入 144 步", 7.55, 3.28, 2.0, 0.25, size=10, color=INK, bold=True)
    add_text(s, "输出 24 步", 9.35, 2.48, 2.0, 0.25, size=10, color=ACCENT, bold=True)
    add_text(s, "可调", 11.3, 2.48, 0.8, 0.22, size=8.5, color=MUTED)
    add_text(s, "递归多步：每次预测回填输入窗口，生成 H 步轨迹。", 0.9, 3.56, 6.2, 0.28, size=9, color=MUTED)

    add_line(s, 0.83, 4.22, 12.5, 4.22, LINE, 0.6)
    groups = [
        ("基线", ["朴素", "季节朴素", "移动平均"], 0.83, 2.2),
        ("统计", ["指数平滑", "Holt-Winters", "SARIMA"], 3.28, 2.2),
        ("机器学习", ["Ridge · Lasso · KNN", "随机森林 · 极端随机树", "GBDT · HistGB · SVR"], 5.95, 4.5),
        ("深度", ["MLP"], 10.85, 1.5),
    ]
    for tag, names, x, w in groups:
        add_text(s, tag, x, 4.42, w, 0.22, size=8.5, color=MUTED)
        add_paragraphs(s, names, x, 4.7, w, 1.4, size=9.5, color=INK, line_spacing=1.35, space_after=2)

    add_text(s, "15", 0.83, 6.28, 0.6, 0.4, size=24, color=ACCENT, bold=True)
    add_text(s, "种算法，统一注册表驱动参数表单；训练在后台执行，逐轮回报损失与日志。", 1.45, 6.38, 10.2, 0.3, size=9.5, color=MUTED)
    add_notes(s, "45 秒：预测先定窗口，训练历史、输入、输出都能调。下面是 15 个算法，分基线、统计、机器学习和深度四类。参数表单由注册表生成，训练任务后台跑，能看到损失曲线和日志。")

    # 6 result
    s = new_slide()
    header(s, "RESULT", "结果：真实值与预测值对齐", "米兰 DUOMO，24 步回测", 6)
    add_text(s, "10 分钟粒度，预测 4 小时", 0.83, 1.7, 4.2, 0.22, size=8.5, color=MUTED)
    horizon = gb["horizon"]
    labels = []
    for i, h in enumerate(horizon):
        labels.append(h["time"][11:16] if i % 6 == 0 else "")
    make_chart(
        s,
        0.83,
        1.95,
        7.1,
        3.4,
        labels,
        [
            ("真实值", [h["actual"] for h in horizon], INK, False),
            ("预测值", [h["predicted"] for h in horizon], ACCENT, True),
        ],
        y_min=9000,
        y_max=13500,
        y_unit=1500,
        legend=True,
    )

    metrics = [
        ("MAE", "449", "平均绝对误差"),
        ("MAPE", "4.49%", "平均绝对百分比误差"),
        ("R²", "0.667", "决定系数"),
        ("Skill", "93.6%", "相对季节朴素基线"),
    ]
    for i, (label, value, note) in enumerate(metrics):
        x = 8.15 + (i % 2) * 2.2
        y = 1.92 + (i // 2) * 1.38
        add_text(s, label, x, y, 1.0, 0.2, size=8.5, color=MUTED)
        add_text(s, value, x, y + 0.2, 2.0, 0.45, size=26, color=INK if i < 2 else ACCENT, bold=True, line_spacing=1.0)
        add_text(s, note, x, y + 0.78, 2.0, 0.2, size=8, color=DIM)

    table_data = [
        ["模型", "MAE", "RMSE", "MAPE", "R²", "Skill"],
        ["GBDT", "449", "578", "4.49%", "0.667", "93.6%"],
        ["随机森林", "2,342", "2,403", "22.59%", "-4.77", "73.5%"],
    ]
    add_table(s, table_data, 0.83, 5.55, [1.9, 1.6, 1.6, 1.6, 1.5, 1.5], row_h=0.34, font_size=9)
    add_text(s, "示例窗口跨元旦，R² 对窗口敏感；平台保留真实结果，不做修饰。", 0.83, 6.55, 9.2, 0.25, size=8.5, color=MUTED)
    add_text(s, "来源：backend/runs.db", 9.95, 7.06, 1.8, 0.18, size=7.5, color=DIM, align=PP_ALIGN.RIGHT)
    add_notes(s, "45 秒：这是 DUOMO 的 24 步回测。GBDT 的 MAE 是 449，MAPE 4.49%，相对季节朴素基线 Skill 93.6%；随机森林明显更差。这个窗口跨元旦，所以 R² 会受节日影响，平台没有修饰结果。")

    # 7 conclusion
    s = new_slide()
    header(s, "CLOSING", "结论：能跑，也能解释", None, 7)
    cols = [
        ("01  已完成", ["五个模块全部打通", "三套数据，统一接口", "15 种算法，参数可调"]),
        ("02  已知边界", ["SARIMA 季节项较慢", "预测区间为近似值", "跨数据集单位不可直接比"]),
        ("03  下一步", ["报告导出与分享", "实时流数据接入", "多变量与空间模型"]),
    ]
    for i, (tag, items) in enumerate(cols):
        x = 0.83 + i * 3.95
        add_text(s, tag, x, 2.0, 3.4, 0.25, size=8.5, color=MUTED)
        add_paragraphs(
            s,
            [{"runs": [{"text": "·  ", "color": ACCENT, "bold": True}, {"text": t}]} for t in items],
            x,
            2.45,
            3.45,
            2.2,
            size=12,
            color=INK,
            line_spacing=1.45,
            space_after=10,
        )
    add_line(s, 0.83, 5.35, 12.5, 5.35, LINE, 0.6)
    add_text(s, "从探索到预测，形成可复现的闭环。", 0.83, 5.7, 9.2, 0.45, size=19, color=INK, bold=True)
    add_text(s, "Q & A", 11.25, 5.78, 1.3, 0.3, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    add_notes(s, "25 秒：最后总结。五个模块打通，三套数据统一接口，15 种算法可调。已知边界主要是 SARIMA 季节项慢、区间是近似、跨数据集单位不能直接比。下一步做导出、实时流和多变量模型。")

    prs.save(OUT)
    print(f"saved {OUT} ({len(prs.slides)} slides)")

